"""Asset inventory helpers: Excel/CSV headers, parse, serialize."""

from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import Workbook, load_workbook

# Exact Excel headers from latest compliteinv.xlsx
EXCEL_HEADERS = [
    "IC No.",
    "Item Name",
    "Item Description",
    "Net Effective Value (Rs.)",
    "Inventory Type",
    "Inventory Location.",
    "GIS Date",
    "Inv Type",
    "remark",
]

FIELD_BY_NORM = {
    "ic no": "ic_no",
    "ic_no": "ic_no",
    "item name": "item_name",
    "item_name": "item_name",
    "item description": "item_description",
    "item_description": "item_description",
    "net effective value (rs)": "net_effective_value",
    "net effective value rs": "net_effective_value",
    "net_effective_value": "net_effective_value",
    "inventory type": "inventory_type",
    "inventory_type": "inventory_type",
    "inventory location": "inventory_location",
    "inventory_location": "inventory_location",
    "gis date": "gis_date",
    "gis_date": "gis_date",
    "inv type": "inv_type",
    "inv_type": "inv_type",
    "remark": "remark",
    "remarks": "remark",
}

MUTABLE_FIELDS = (
    "ic_no",
    "item_name",
    "item_description",
    "net_effective_value",
    "inventory_type",
    "inventory_location",
    "gis_date",
    "inv_type",
    "remark",
    "assigned_to",
)


def is_assigned_item(item: Any) -> bool:
    """Assigned when issued to someone, or Excel Inventory Type is Personal."""
    assigned_to = str(getattr(item, "assigned_to", "") or "").strip()
    if assigned_to:
        return True
    inv_type = str(getattr(item, "inventory_type", "") or "").strip().lower()
    return inv_type == "personal"


def normalize_header(raw: Any) -> str:
    """Normalize Excel headers for mapping.

    Strips all dots so ``Net Effective Value (Rs.)`` and ``Inventory Location.``
    both match the canonical FIELD_BY_NORM keys (``… (rs)``, ``inventory location``).
    """
    text = str(raw or "").strip().lower()
    text = text.replace(".", "")
    text = re.sub(r"\s+", " ", text)
    return text


def map_headers(headers: list[Any]) -> dict[int, str]:
    mapping: dict[int, str] = {}
    for idx, h in enumerate(headers):
        key = FIELD_BY_NORM.get(normalize_header(h))
        if key:
            mapping[idx] = key
    return mapping


def parse_gis_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    # Excel serial as string
    try:
        serial = float(text)
        # openpyxl sometimes already gives date; if numeric days since 1899-12-30
        from datetime import timedelta

        base = datetime(1899, 12, 30)
        return (base + timedelta(days=serial)).date()
    except (TypeError, ValueError, OverflowError):
        return None


def parse_value(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        return float(Decimal(text))
    except (InvalidOperation, ValueError):
        return None


def row_dict_from_cells(mapping: dict[int, str], cells: list[Any]) -> dict[str, Any]:
    out: dict[str, Any] = {f: "" for f in MUTABLE_FIELDS if f not in {"net_effective_value", "gis_date"}}
    out["net_effective_value"] = None
    out["gis_date"] = None
    for idx, field in mapping.items():
        if idx >= len(cells):
            continue
        raw = cells[idx]
        if field == "net_effective_value":
            out[field] = parse_value(raw)
        elif field == "gis_date":
            out[field] = parse_gis_date(raw)
        else:
            out[field] = str(raw or "").strip()
    out["ic_no"] = str(out.get("ic_no") or "").strip()
    return out


def item_to_excel_row(item: Any) -> list[Any]:
    gis = getattr(item, "gis_date", None)
    gis_out = gis.strftime("%d-%m-%Y") if gis else ""
    val = getattr(item, "net_effective_value", None)
    return [
        getattr(item, "ic_no", "") or "",
        getattr(item, "item_name", "") or "",
        getattr(item, "item_description", "") or "",
        val if val is not None else "",
        getattr(item, "inventory_type", "") or "",
        getattr(item, "inventory_location", "") or "",
        gis_out,
        getattr(item, "inv_type", "") or "",
        getattr(item, "remark", "") or "",
    ]


# Example first data row so users see the expected format (delete before real import).
TEMPLATE_EXAMPLE_ROW = [
    "IRD/3371/2026/1",
    "Dell Precision 7960 Workstation",
    "Example description — replace this row with real data",
    "1797767.00",
    "Lab",
    "ARCHITECTURE LAB",
    "17-03-2026",
    "IRD",
    "sample row — delete before import",
]


def build_template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"
    ws.append(EXCEL_HEADERS)
    ws.append(TEMPLATE_EXAMPLE_ROW)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_template_csv() -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(EXCEL_HEADERS)
    writer.writerow(TEMPLATE_EXAMPLE_ROW)
    return buf.getvalue().encode("utf-8-sig")


def build_export_xlsx(items: list[Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"
    ws.append(EXCEL_HEADERS)
    for item in items:
        ws.append(item_to_excel_row(item))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_export_csv(items: list[Any]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(EXCEL_HEADERS)
    for item in items:
        writer.writerow(item_to_excel_row(item))
    return buf.getvalue().encode("utf-8-sig")


def parse_import_bytes(data: bytes, filename: str = "") -> list[dict[str, Any]]:
    name = (filename or "").lower()
    if name.endswith(".csv") or (not name.endswith(".xlsx") and b"," in data[:200]):
        return _parse_csv(data)
    return _parse_xlsx(data)


def _parse_xlsx(data: bytes) -> list[dict[str, Any]]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = list(next(rows_iter))
    except StopIteration:
        wb.close()
        return []
    mapping = map_headers(header)
    if "ic_no" not in mapping.values():
        wb.close()
        raise ValueError("Import file must include an 'IC No.' column")
    out: list[dict[str, Any]] = []
    for row in rows_iter:
        cells = list(row)
        parsed = row_dict_from_cells(mapping, cells)
        if not parsed["ic_no"]:
            continue
        out.append(parsed)
    wb.close()
    return out


def _parse_csv(data: bytes) -> list[dict[str, Any]]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    try:
        header = next(reader)
    except StopIteration:
        return []
    mapping = map_headers(header)
    if "ic_no" not in mapping.values():
        raise ValueError("Import file must include an 'IC No.' column")
    out: list[dict[str, Any]] = []
    for row in reader:
        parsed = row_dict_from_cells(mapping, list(row))
        if not parsed["ic_no"]:
            continue
        out.append(parsed)
    return out
