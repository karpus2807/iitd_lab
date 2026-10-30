"""Scholar enrolment Excel/CSV template, parse, export."""

from __future__ import annotations

import csv
import io
import re
from typing import Any

from openpyxl import Workbook, load_workbook

EXCEL_HEADERS = [
    "Lab Name",
    "Name",
    "ID",
    "Group",
    "Faculty",
    "System Required",
    "Inventory IC/Label",
    "Email",
    "Phone",
    "Notes",
]

TEMPLATE_EXAMPLE_ROW = [
    "ARCHITECTURE LAB",
    "Ada Example",
    "2026CSZ0001",
    "Ph.D.",
    "Prof. Example",
    "Yes",
    "IRD/3371/2026/1",
    "ada.example@iitd.ac.in",
    "9876543210",
    "sample row — delete before real import",
]

FIELD_BY_NORM = {
    "lab name": "lab_name",
    "lab": "lab_name",
    "name": "full_name",
    "full name": "full_name",
    "id": "scholar_id",
    "scholar id": "scholar_id",
    "entry id": "scholar_id",
    "roll no": "scholar_id",
    "roll number": "scholar_id",
    "group": "group_name",
    "group name": "group_name",
    "faculty": "faculty",
    "supervisor": "faculty",
    "system required": "system_required",
    "system": "system_required",
    "inventory ic/label": "inventory_label",
    "inventory ic label": "inventory_label",
    "inventory ic": "inventory_label",
    "inventory label": "inventory_label",
    "inventory": "inventory_label",
    "email": "email",
    "phone": "phone",
    "contact": "phone",
    "contact number": "phone",
    "notes": "notes",
    "remark": "notes",
    "remarks": "notes",
}


def normalize_header(raw: Any) -> str:
    text = str(raw or "").strip().lower()
    text = text.replace(".", "")
    text = re.sub(r"\s+", " ", text)
    return text


def map_headers(headers: list[Any]) -> dict[int, str]:
    mapping: dict[int, str] = {}
    for idx, h in enumerate(headers):
        key = FIELD_BY_NORM.get(normalize_header(h))
        if key:
            mapping[idx] = key
    return mapping


def parse_yes_no(value: Any) -> bool:
    text = str(value or "").strip().lower()
    return text in {"1", "true", "yes", "y"}


def row_dict_from_cells(mapping: dict[int, str], cells: list[Any]) -> dict[str, Any]:
    out: dict[str, Any] = {
        "lab_name": "",
        "full_name": "",
        "scholar_id": "",
        "group_name": "",
        "faculty": "",
        "system_required": False,
        "inventory_label": "",
        "email": "",
        "phone": "",
        "notes": "",
    }
    for idx, field in mapping.items():
        if idx >= len(cells):
            continue
        raw = cells[idx]
        if field == "system_required":
            out[field] = parse_yes_no(raw)
        else:
            out[field] = str(raw or "").strip()
    out["full_name"] = str(out.get("full_name") or "").strip()
    out["lab_name"] = str(out.get("lab_name") or "").strip()
    return out


def is_example_row(parsed: dict[str, Any]) -> bool:
    notes = (parsed.get("notes") or "").lower()
    name = (parsed.get("full_name") or "").lower()
    if "sample row" in notes or "delete before" in notes:
        return True
    if name.endswith(" example") or name == "ada example":
        return True
    return False


def scholar_to_excel_row(item: Any, lab_name: str = "") -> list[Any]:
    return [
        lab_name or getattr(item, "lab_name", "") or "",
        getattr(item, "full_name", "") or "",
        getattr(item, "scholar_id", "") or "",
        getattr(item, "group_name", "") or "",
        getattr(item, "faculty", "") or "",
        "Yes" if getattr(item, "system_required", False) else "No",
        getattr(item, "inventory_label", "") or "",
        getattr(item, "email", "") or "",
        getattr(item, "phone", "") or "",
        getattr(item, "notes", "") or "",
    ]


def build_template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Scholars"
    ws.append(EXCEL_HEADERS)
    ws.append(TEMPLATE_EXAMPLE_ROW)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_template_csv() -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(EXCEL_HEADERS)
    writer.writerow(TEMPLATE_EXAMPLE_ROW)
    return buf.getvalue().encode("utf-8-sig")


def build_export_xlsx(rows: list[tuple[Any, str]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Scholars"
    ws.append(EXCEL_HEADERS)
    for item, lab_name in rows:
        ws.append(scholar_to_excel_row(item, lab_name))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_export_csv(rows: list[tuple[Any, str]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(EXCEL_HEADERS)
    for item, lab_name in rows:
        writer.writerow(scholar_to_excel_row(item, lab_name))
    return buf.getvalue().encode("utf-8-sig")


def parse_import_bytes(data: bytes, filename: str = "") -> list[dict[str, Any]]:
    name = (filename or "").lower()
    if name.endswith(".csv") or (not name.endswith(".xlsx") and b"," in data[:200]):
        return _parse_csv(data)
    return _parse_xlsx(data)


def _parse_xlsx(data: bytes) -> list[dict[str, Any]]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = list(next(rows_iter))
    except StopIteration:
        wb.close()
        return []
    mapping = map_headers(header)
    if "full_name" not in mapping.values() or "lab_name" not in mapping.values():
        wb.close()
        raise ValueError("Import file must include 'Lab Name' and 'Name' columns")
    out: list[dict[str, Any]] = []
    for row in rows_iter:
        parsed = row_dict_from_cells(mapping, list(row))
        if not parsed["full_name"] or is_example_row(parsed):
            continue
        out.append(parsed)
    wb.close()
    return out


def _parse_csv(data: bytes) -> list[dict[str, Any]]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    try:
        header = next(reader)
    except StopIteration:
        return []
    mapping = map_headers(header)
    if "full_name" not in mapping.values() or "lab_name" not in mapping.values():
        raise ValueError("Import file must include 'Lab Name' and 'Name' columns")
    out: list[dict[str, Any]] = []
    for row in reader:
        parsed = row_dict_from_cells(mapping, list(row))
        if not parsed["full_name"] or is_example_row(parsed):
            continue
        out.append(parsed)
    return out
