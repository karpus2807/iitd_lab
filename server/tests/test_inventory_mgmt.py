"""Asset inventory registry (Inv Management) API."""

import io

import pytest
from httpx import AsyncClient
from openpyxl import Workbook, load_workbook

from app.services.inv_mgmt import EXCEL_HEADERS, map_headers, normalize_header, parse_import_bytes


async def _create(client: AsyncClient, headers: dict, ic: str, **extra):
    body = {
        "ic_no": ic,
        "item_name": extra.get("item_name", "Workstation"),
        "item_description": extra.get("item_description", "desc"),
        "net_effective_value": extra.get("net_effective_value", 1000.5),
        "inventory_type": extra.get("inventory_type", "Lab"),
        "inventory_location": extra.get("inventory_location", "CLOUD LAB"),
        "gis_date": extra.get("gis_date", "2026-03-14"),
        "inv_type": extra.get("inv_type", "IRD"),
        "remark": extra.get("remark", ""),
    }
    r = await client.post("/api/inventory", headers=headers, json=body)
    assert r.status_code == 200, r.text
    return r.json()


@pytest.mark.asyncio
async def test_crud_remove_and_list(client: AsyncClient, auth_headers: dict):
    a = await _create(client, auth_headers, "IRD/TEST/2026/1")
    listed = await client.get("/api/inventory?status=active&q=IRD/TEST", headers=auth_headers)
    assert listed.status_code == 200
    body = listed.json()
    assert body["total"] >= 1
    assert any(x["ic_no"] == "IRD/TEST/2026/1" for x in body["items"])

    patched = await client.patch(
        f"/api/inventory/{a['id']}",
        headers=auth_headers,
        json={"item_name": "Updated WS", "inventory_location": "ARCH LAB"},
    )
    assert patched.status_code == 200
    assert patched.json()["item_name"] == "Updated WS"

    rem = await client.post(
        "/api/inventory/remove",
        headers=auth_headers,
        json={"ids": [a["id"]], "remark": "Transferred out of lab"},
    )
    assert rem.status_code == 200, rem.text
    assert rem.json()["removed"] == 1

    removed = await client.get("/api/inventory?status=removed&q=IRD/TEST", headers=auth_headers)
    assert removed.status_code == 200
    row = next(x for x in removed.json()["items"] if x["id"] == a["id"])
    assert row["status"] == "removed"
    assert row["removal_remark"] == "Transferred out of lab"
    assert row["removed_by"] == "admin"


@pytest.mark.asyncio
async def test_restore_and_purge(client: AsyncClient, auth_headers: dict):
    a = await _create(client, auth_headers, "IRD/TEST/2026/RESTORE")
    b = await _create(client, auth_headers, "IRD/TEST/2026/PURGE")
    rem = await client.post(
        "/api/inventory/remove",
        headers=auth_headers,
        json={"ids": [a["id"], b["id"]], "remark": "temp out"},
    )
    assert rem.status_code == 200
    assert rem.json()["removed"] == 2

    restored = await client.post(
        "/api/inventory/restore",
        headers=auth_headers,
        json={"ids": [a["id"]]},
    )
    assert restored.status_code == 200, restored.text
    assert restored.json()["restored"] == 1
    active = await client.get("/api/inventory?status=active&q=IRD/TEST/2026/RESTORE", headers=auth_headers)
    row = next(x for x in active.json()["items"] if x["id"] == a["id"])
    assert row["status"] == "active"
    assert row["removal_remark"] is None
    assert row["removed_at"] is None

    # Active items are skipped by purge
    skip = await client.post(
        "/api/inventory/purge",
        headers=auth_headers,
        json={"ids": [a["id"]]},
    )
    assert skip.status_code == 200
    assert skip.json()["deleted"] == 0
    assert a["ic_no"] in skip.json()["skipped"]

    purged = await client.post(
        "/api/inventory/purge",
        headers=auth_headers,
        json={"ids": [b["id"]]},
    )
    assert purged.status_code == 200, purged.text
    assert purged.json()["deleted"] == 1
    gone = await client.get(f"/api/inventory/{b['id']}", headers=auth_headers)
    assert gone.status_code == 404


@pytest.mark.asyncio
async def test_inventory_filters_and_facets(client: AsyncClient, auth_headers: dict):
    await _create(
        client,
        auth_headers,
        "IRD/FILT/LAB",
        inventory_type="Lab",
        inventory_location="VLSI LAB",
        inv_type="Institute",
    )
    await _create(
        client,
        auth_headers,
        "IRD/FILT/PERSONAL",
        inventory_type="Personal",
        inventory_location="CSE",
        inv_type="IRD",
    )
    facets = await client.get("/api/inventory/facets?status=active", headers=auth_headers)
    assert facets.status_code == 200, facets.text
    body = facets.json()
    assert "Lab" in body["inventory_types"]
    assert "Personal" in body["inventory_types"]
    assert "VLSI LAB" in body["inventory_locations"]
    assert "IRD" in body["inv_types"]

    only_lab = await client.get(
        "/api/inventory?status=active&inventory_type=Lab&q=IRD/FILT",
        headers=auth_headers,
    )
    assert only_lab.status_code == 200
    ics = {x["ic_no"] for x in only_lab.json()["items"]}
    assert "IRD/FILT/LAB" in ics
    assert "IRD/FILT/PERSONAL" not in ics

    only_ird = await client.get(
        "/api/inventory?status=active&inv_type=IRD&q=IRD/FILT",
        headers=auth_headers,
    )
    assert {x["ic_no"] for x in only_ird.json()["items"]} == {"IRD/FILT/PERSONAL"}

    assigned = await client.get(
        "/api/inventory?status=active&link=assigned&q=IRD/FILT",
        headers=auth_headers,
    )
    assert {x["ic_no"] for x in assigned.json()["items"]} == {"IRD/FILT/PERSONAL"}


@pytest.mark.asyncio
async def test_empty_removed(client: AsyncClient, auth_headers: dict):
    a = await _create(client, auth_headers, "IRD/TEST/2026/EMPTY-A")
    b = await _create(client, auth_headers, "IRD/TEST/2026/EMPTY-B")
    rem = await client.post(
        "/api/inventory/remove",
        headers=auth_headers,
        json={"ids": [a["id"], b["id"]], "remark": "bulk wipe"},
    )
    assert rem.status_code == 200
    emptied = await client.post("/api/inventory/empty-removed", headers=auth_headers)
    assert emptied.status_code == 200, emptied.text
    assert emptied.json()["deleted"] >= 2
    listed = await client.get("/api/inventory?status=removed&q=IRD/TEST/2026/EMPTY", headers=auth_headers)
    assert listed.status_code == 200
    assert listed.json()["total"] == 0
    gone_a = await client.get(f"/api/inventory/{a['id']}", headers=auth_headers)
    gone_b = await client.get(f"/api/inventory/{b['id']}", headers=auth_headers)
    assert gone_a.status_code == 404
    assert gone_b.status_code == 404


@pytest.mark.asyncio
async def test_remove_requires_remark(client: AsyncClient, auth_headers: dict):
    a = await _create(client, auth_headers, "IRD/TEST/2026/2")
    bad = await client.post(
        "/api/inventory/remove",
        headers=auth_headers,
        json={"ids": [a["id"]], "remark": ""},
    )
    assert bad.status_code == 422


def test_excel_header_maps_net_value_with_rs_dot():
    """Excel uses 'Net Effective Value (Rs.)' — period inside parens must still map."""
    headers = [
        "IC No.",
        "Item Name",
        "Item Description",
        "Net Effective Value (Rs.)",
        "Inventory Type",
        "Inventory Location.",
        "GIS Date",
        "Inv Type",
        "remark",
    ]
    assert normalize_header("Net Effective Value (Rs.)") == "net effective value (rs)"
    mapping = map_headers(headers)
    assert set(mapping.values()) == {
        "ic_no",
        "item_name",
        "item_description",
        "net_effective_value",
        "inventory_type",
        "inventory_location",
        "gis_date",
        "inv_type",
        "remark",
    }

    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    ws.append(
        [
            "IRD/VAL/1",
            "Box",
            "desc",
            "1797767.00",
            "Personal",
            "CSE",
            "17-03-2026",
            "IRD",
            "",
        ]
    )
    buf = io.BytesIO()
    wb.save(buf)
    parsed = parse_import_bytes(buf.getvalue(), "sample.xlsx")
    assert len(parsed) == 1
    assert parsed[0]["net_effective_value"] == 1797767.0
    assert parsed[0]["inventory_type"] == "Personal"
    assert parsed[0]["inventory_location"] == "CSE"


@pytest.mark.asyncio
async def test_registered_and_assigned_status(client: AsyncClient, auth_headers: dict):
    await _create(
        client,
        auth_headers,
        "IRD/LINK/PERSONAL",
        item_name="Laptop",
        inventory_type="Personal",
    )
    lab_item = await _create(
        client,
        auth_headers,
        "IRD/LINK/LAB",
        item_name="Server",
        inventory_type="",
        inventory_location="",
    )
    from app.db import get_session_factory
    from app.models import Lab, Machine
    from sqlalchemy import select

    factory = get_session_factory()
    async with factory() as db:
        lab = (
            await db.execute(select(Lab).where(Lab.name == "ARCHITECTURE LAB"))
        ).scalar_one_or_none()
        if not lab:
            lab = Lab(name="ARCHITECTURE LAB", code="ARCH")
            db.add(lab)
            await db.flush()
        db.add(
            Machine(
                hostname="link-host",
                inventory_id="IRD/LINK/LAB",
                os_name="Linux",
                approved=True,
                lab_id=lab.id,
            )
        )
        await db.commit()

    listed = await client.get("/api/inventory?q=IRD/LINK", headers=auth_headers)
    assert listed.status_code == 200
    by_ic = {x["ic_no"]: x for x in listed.json()["items"]}
    assert by_ic["IRD/LINK/PERSONAL"]["assigned"] is True
    assert by_ic["IRD/LINK/PERSONAL"]["registered"] is False
    assert by_ic["IRD/LINK/LAB"]["registered"] is True
    assert by_ic["IRD/LINK/LAB"]["assigned"] is False
    assert by_ic["IRD/LINK/LAB"]["machine_hostname"] == "link-host"
    # Blank template location/type → fall back to machine lab / Lab
    assert by_ic["IRD/LINK/LAB"]["inventory_location"] == "ARCHITECTURE LAB"
    assert by_ic["IRD/LINK/LAB"]["inventory_type"] == "Lab"
    assert by_ic["IRD/LINK/LAB"]["machine_lab_name"] == "ARCHITECTURE LAB"

    patched = await client.patch(
        f"/api/inventory/{lab_item['id']}",
        headers=auth_headers,
        json={"assigned_to": "Student A"},
    )
    assert patched.status_code == 200
    assert patched.json()["assigned"] is True
    assert patched.json()["registered"] is True
    assert patched.json()["assigned_to"] == "Student A"
    assert patched.json()["inventory_location"] == "ARCHITECTURE LAB"


@pytest.mark.asyncio
async def test_template_headers_and_import_upsert(client: AsyncClient, auth_headers: dict):
    tmpl = await client.get("/api/inventory/template.xlsx", headers=auth_headers)
    assert tmpl.status_code == 200
    wb = load_workbook(io.BytesIO(tmpl.content), read_only=True)
    headers = [c.value for c in next(wb.active.iter_rows(max_row=1))]
    wb.close()
    assert headers == EXCEL_HEADERS

    out = Workbook()
    ws = out.active
    ws.append(EXCEL_HEADERS)
    ws.append(
        [
            "IRD/IMP/2026/1",
            "GPU Box",
            "Dell tower",
            50000,
            "Lab",
            "VLSI LAB",
            "01-02-2026",
            "Institute",
            "ok",
        ]
    )
    buf = io.BytesIO()
    out.save(buf)
    files = {
        "file": (
            "import.xlsx",
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    imp = await client.post("/api/inventory/import", headers=auth_headers, files=files)
    assert imp.status_code == 200, imp.text
    assert imp.json()["created"] == 1

    # Upsert same IC
    out2 = Workbook()
    ws2 = out2.active
    ws2.append(EXCEL_HEADERS)
    ws2.append(
        [
            "IRD/IMP/2026/1",
            "GPU Box v2",
            "Dell tower",
            51000,
            "Lab",
            "VLSI LAB",
            "01-02-2026",
            "Institute",
            "updated",
        ]
    )
    buf2 = io.BytesIO()
    out2.save(buf2)
    imp2 = await client.post(
        "/api/inventory/import",
        headers=auth_headers,
        files={
            "file": (
                "import2.xlsx",
                buf2.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert imp2.status_code == 200
    assert imp2.json()["updated"] == 1

    listed = await client.get("/api/inventory?q=IRD/IMP/2026/1", headers=auth_headers)
    item = next(x for x in listed.json()["items"] if x["ic_no"] == "IRD/IMP/2026/1")
    assert item["item_name"] == "GPU Box v2"
    assert item["remark"] == "updated"


@pytest.mark.asyncio
async def test_suggest_and_batch(client: AsyncClient, auth_headers: dict):
    a = await _create(client, auth_headers, "IRD/SUG/1", item_name="Alpha Kit")
    b = await _create(client, auth_headers, "IRD/SUG/2", item_name="Beta Kit")
    sug = await client.get("/api/inventory/suggest?q=IRD/SUG", headers=auth_headers)
    assert sug.status_code == 200
    assert len(sug.json()) >= 2

    batch = await client.post(
        "/api/inventory/batch",
        headers=auth_headers,
        json={"ids": [a["id"], b["id"]], "fields": {"inventory_location": "BATCH LAB"}},
    )
    assert batch.status_code == 200
    assert batch.json()["updated"] == 2
    listed = await client.get("/api/inventory?q=IRD/SUG", headers=auth_headers)
    locs = {x["ic_no"]: x["inventory_location"] for x in listed.json()["items"] if x["ic_no"].startswith("IRD/SUG")}
    assert locs["IRD/SUG/1"] == "BATCH LAB"
    assert locs["IRD/SUG/2"] == "BATCH LAB"
