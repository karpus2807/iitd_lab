"""Scholars enrolment API."""

import io
from unittest.mock import AsyncMock

import pytest
from httpx import AsyncClient
from openpyxl import Workbook, load_workbook

from app.services.scholar_mgmt import EXCEL_HEADERS, TEMPLATE_EXAMPLE_ROW


async def _enroll_with_otp(client: AsyncClient, headers: dict, data: dict, code: str = "111111") -> dict:
    """Inventory assign / OTP enrol returns pending_otp until verify-otp."""
    pending = await client.post("/api/scholars", headers=headers, data=data)
    assert pending.status_code == 200, pending.text
    body = pending.json()
    if not body.get("pending_otp"):
        return body
    verified = await client.post(
        "/api/scholars/verify-otp",
        headers=headers,
        data={"challenge_id": body["challenge_id"], "code": code},
    )
    assert verified.status_code == 200, verified.text
    return verified.json()["scholar"]


@pytest.mark.asyncio
async def test_enroll_listed_and_non_list(client: AsyncClient, auth_headers: dict, monkeypatch):
    monkeypatch.setattr("app.services.email_otp.send_email", AsyncMock())
    monkeypatch.setattr("app.services.email_otp.generate_code", lambda: "111111")
    await client.put(
        "/api/admin/smtp",
        headers=auth_headers,
        json={
            "enabled": True,
            "host": "smtp.test.local",
            "port": 587,
            "from_addr": "lab@test.local",
            "password": "x",
        },
    )

    created_lab = await client.post(
        "/api/labs",
        headers=auth_headers,
        json={"name": "Scholar Test Lab", "room": "401", "building": "Bharti"},
    )
    assert created_lab.status_code == 200, created_lab.text
    lab = created_lab.json()
    assert lab["room"] == "401"

    inv = await client.post(
        "/api/inventory",
        headers=auth_headers,
        json={
            "ic_no": "IRD/SCH/1",
            "item_name": "Scholar PC",
            "inventory_type": "Lab",
            "inventory_location": lab["name"],
            "inv_type": "Institute",
        },
    )
    assert inv.status_code == 200, inv.text
    inv_id = inv.json()["id"]

    fd = {
        "lab_id": str(lab["id"]),
        "full_name": "Ada Scholar",
        "scholar_id": "2026CSZ0001",
        "group_name": "Ph.D.",
        "faculty": "Prof X",
        "system_required": "true",
        "inventory_kind": "listed",
        "inventory_item_id": inv_id,
        "email": "ada@example.com",
        "phone": "9999999999",
    }
    body = await _enroll_with_otp(client, auth_headers, fd)
    assert body["full_name"] == "Ada Scholar"
    assert body["lab_name"] == lab["name"]
    assert body["inventory_kind"] == "listed"
    assert body["inventory_label"] == "IRD/SCH/1"
    assert body["system_required"] is True

    linked = await client.get(f"/api/inventory/{inv_id}", headers=auth_headers)
    assert linked.json()["assigned_to"] == "Ada Scholar"
    assert linked.json()["assigned"] is True

    fd2 = {
        "lab_id": str(lab["id"]),
        "full_name": "Postdoc Bob",
        "scholar_id": "PDF-9",
        "group_name": "Postdoc",
        "faculty": "Prof Y",
        "system_required": "true",
        "inventory_kind": "non_list",
        "inventory_label": "Personal MacBook",
        "email": "bob@example.com",
    }
    created2 = await _enroll_with_otp(client, auth_headers, fd2)
    assert created2["inventory_kind"] == "non_list"

    fd3 = {
        "lab_id": str(lab["id"]),
        "full_name": "Intern Cara",
        "group_name": "Intern",
        "system_required": "false",
        "inventory_kind": "none",
    }
    created3 = await client.post("/api/scholars", headers=auth_headers, data=fd3)
    assert created3.status_code == 200, created3.text
    assert created3.json()["system_required"] is False

    listed = await client.get(f"/api/scholars?lab_id={lab['id']}", headers=auth_headers)
    assert listed.status_code == 200
    assert listed.json()["total"] >= 3

    facets = await client.get("/api/scholars/facets", headers=auth_headers)
    assert facets.status_code == 200
    assert "Ph.D." in facets.json()["groups"]


@pytest.mark.asyncio
async def test_scholar_template_import_remove_restore_purge(client: AsyncClient, auth_headers: dict):
    lab = (
        await client.post(
            "/api/labs",
            headers=auth_headers,
            json={"name": "Import Lab", "room": "101"},
        )
    ).json()

    tmpl = await client.get("/api/scholars/template.xlsx", headers=auth_headers)
    assert tmpl.status_code == 200
    wb = load_workbook(io.BytesIO(tmpl.content), read_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    wb.close()
    assert list(rows[0]) == EXCEL_HEADERS
    assert list(rows[1])[:3] == TEMPLATE_EXAMPLE_ROW[:3]

    out = Workbook()
    ws = out.active
    ws.append(EXCEL_HEADERS)
    ws.append(TEMPLATE_EXAMPLE_ROW)  # skipped as example
    ws.append(
        [
            "Import Lab",
            "7",
            "Real Scholar",
            "IMP-1",
            "M.Tech",
            "Prof Z",
            "No",
            "",
            "real@example.com",
            "111",
            "No",
            "real row",
        ]
    )
    buf = io.BytesIO()
    out.save(buf)
    imp = await client.post(
        "/api/scholars/import",
        headers=auth_headers,
        files={
            "file": (
                "scholars.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert imp.status_code == 200, imp.text
    assert imp.json()["created"] == 1

    listed = await client.get(
        f"/api/scholars?lab_id={lab['id']}&q=Real Scholar",
        headers=auth_headers,
    )
    item = next(x for x in listed.json()["items"] if x["full_name"] == "Real Scholar")
    assert item["bio_id"] == "7"
    assert item["approved"] is True
    assert item["no_dues"] is False

    rem = await client.post(
        "/api/scholars/remove",
        headers=auth_headers,
        json={"ids": [item["id"]], "remark": "left lab"},
    )
    assert rem.status_code == 200
    assert rem.json()["removed"] == 1

    removed = await client.get("/api/scholars?status=removed&q=Real Scholar", headers=auth_headers)
    assert removed.json()["total"] >= 1
    assert removed.json()["items"][0]["removal_remark"] == "left lab"

    restored = await client.post(
        "/api/scholars/restore",
        headers=auth_headers,
        json={"ids": [item["id"]]},
    )
    assert restored.status_code == 200
    assert restored.json()["restored"] == 1

    await client.post(
        "/api/scholars/remove",
        headers=auth_headers,
        json={"ids": [item["id"]], "remark": "wipe"},
    )
    purged = await client.post(
        "/api/scholars/purge",
        headers=auth_headers,
        json={"ids": [item["id"]]},
    )
    assert purged.status_code == 200
    assert purged.json()["deleted"] == 1
    gone = await client.get(f"/api/scholars/{item['id']}", headers=auth_headers)
    assert gone.status_code == 404


@pytest.mark.asyncio
async def test_scholars_dashboard_summary_and_lab(client: AsyncClient, auth_headers: dict):
    lab = (
        await client.post(
            "/api/labs",
            headers=auth_headers,
            json={"name": "Dash Lab", "room": "9"},
        )
    ).json()
    created = await client.post(
        "/api/scholars",
        headers=auth_headers,
        data={
            "lab_id": str(lab["id"]),
            "full_name": "Dash Student",
            "scholar_id": "DASH-1",
            "bio_id": "42",
            "system_required": "true",
            "inventory_kind": "non_list",
            "inventory_label": "Personal PC",
            "verify_method": "document",
            "require_otp": "false",
        },
        files={"document": ("proof.pdf", b"%PDF-1.4\nx\n%%EOF\n", "application/pdf")},
    )
    assert created.status_code == 200, created.text
    assert created.json().get("inventory_assigned_at") or created.json().get("inventory_kind") == "non_list"

    summary = await client.get("/api/scholars/dashboard", headers=auth_headers)
    assert summary.status_code == 200, summary.text
    body = summary.json()
    assert body["scholars"] >= 1
    lab_row = next(x for x in body["labs"] if x["id"] == lab["id"])
    assert lab_row["scholars"] >= 1
    assert lab_row["with_inventory"] >= 1

    detail = await client.get(f"/api/scholars/dashboard/{lab['id']}", headers=auth_headers)
    assert detail.status_code == 200, detail.text
    d = detail.json()
    assert d["lab"]["name"] == "Dash Lab"
    assert any(s["full_name"] == "Dash Student" for s in d["scholars"])
    student = next(s for s in d["scholars"] if s["full_name"] == "Dash Student")
    assert student["inventory"]["ic_no"] == "Personal PC"
    assert student["inventory_assigned_at"]


@pytest.mark.asyncio
async def test_bio_id_per_lab_and_id_school_wide(client: AsyncClient, auth_headers: dict):
    lab_a = (await client.post("/api/labs", headers=auth_headers, json={"name": "VLSI Lab", "room": "1"})).json()
    lab_b = (await client.post("/api/labs", headers=auth_headers, json={"name": "MTech Lab", "room": "2"})).json()

    a1 = await client.post(
        "/api/scholars",
        headers=auth_headers,
        data={
            "lab_id": str(lab_a["id"]),
            "full_name": "Student A",
            "scholar_id": "SCH-GLOBAL-1",
            "bio_id": "1",
            "system_required": "false",
            "inventory_kind": "none",
            "require_otp": "false",
        },
    )
    assert a1.status_code == 200, a1.text
    assert a1.json()["approved"] is False

    # Same bio-id allowed in another lab
    b1 = await client.post(
        "/api/scholars",
        headers=auth_headers,
        data={
            "lab_id": str(lab_b["id"]),
            "full_name": "Student B",
            "scholar_id": "SCH-GLOBAL-2",
            "bio_id": "1",
            "system_required": "false",
            "inventory_kind": "none",
            "require_otp": "false",
        },
    )
    assert b1.status_code == 200, b1.text

    # Same school-wide ID rejected
    clash = await client.post(
        "/api/scholars",
        headers=auth_headers,
        data={
            "lab_id": str(lab_b["id"]),
            "full_name": "Clash",
            "scholar_id": "SCH-GLOBAL-1",
            "bio_id": "99",
            "system_required": "false",
            "inventory_kind": "none",
            "require_otp": "false",
        },
    )
    assert clash.status_code == 400

    # Same bio-id in same lab rejected
    clash_bio = await client.post(
        "/api/scholars",
        headers=auth_headers,
        data={
            "lab_id": str(lab_a["id"]),
            "full_name": "Clash Bio",
            "scholar_id": "SCH-GLOBAL-3",
            "bio_id": "1",
            "system_required": "false",
            "inventory_kind": "none",
            "require_otp": "false",
        },
    )
    assert clash_bio.status_code == 400

    approved = await client.post(f"/api/scholars/{a1.json()['id']}/approve", headers=auth_headers)
    assert approved.status_code == 200
    assert approved.json()["approved"] is True

    dues = await client.post(
        f"/api/scholars/{a1.json()['id']}/no-dues",
        headers=auth_headers,
        data={"no_dues": "true"},
    )
    assert dues.status_code == 200
    assert dues.json()["no_dues"] is True

    revoke = await client.post(
        f"/api/scholars/{a1.json()['id']}/no-dues",
        headers=auth_headers,
        data={"no_dues": "false"},
    )
    assert revoke.status_code == 200
    assert revoke.json()["no_dues"] is False


@pytest.mark.asyncio
async def test_inventory_template_has_example_row(client: AsyncClient, auth_headers: dict):
    from app.services.inv_mgmt import EXCEL_HEADERS as INV_HEADERS
    from app.services.inv_mgmt import TEMPLATE_EXAMPLE_ROW as INV_EXAMPLE

    tmpl = await client.get("/api/inventory/template.xlsx", headers=auth_headers)
    assert tmpl.status_code == 200
    wb = load_workbook(io.BytesIO(tmpl.content), read_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    wb.close()
    assert list(rows[0]) == INV_HEADERS
    assert list(rows[1])[0] == INV_EXAMPLE[0]
